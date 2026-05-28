import os
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from app.models.inconsistencia import Inconsistencia
from app.models.consolidacao import Consolidacao
from app.core.config import settings
from datetime import datetime

AZUL = "1E3A5F"; VERDE = "16A34A"; VERD_CL = "D1FAE5"
AMAR_CL = "FEF3C7"; VERM_CL = "FEE2E2"; CINZA = "F9FAFB"; BORDA = "E5E7EB"
thin = Side(style="thin", color=BORDA)
BDR  = Border(left=thin, right=thin, top=thin, bottom=thin)

def _header(ws, cols_widths, row=1):
    for i, (col, w) in enumerate(cols_widths, 1):
        c = ws.cell(row, i, col)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BDR
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 18

def gerar_base_consolidada(db: Session, competencia_id: int, empresa_nome: str) -> str:
    os.makedirs(settings.EXPORT_DIR, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    
    # aba inconsistencias
    ws_inc = wb.create_sheet("Auditoria")
    cols = [("Módulo",14),("Descrição",55),("Gravidade",12),("Categoria",16),("Status",16),("Tratativa",40)]
    _header(ws_inc, cols)
    incs = db.query(Inconsistencia).filter(Inconsistencia.competencia_id == competencia_id).all()
    grav_fill = {"alta": VERM_CL, "media": AMAR_CL, "baixa": "DBEAFE"}
    for i, inc in enumerate(incs, 2):
        vals = [inc.modulo, inc.descricao, inc.gravidade.value, inc.categoria or "", inc.status.value, inc.tratativa or ""]
        for j, v in enumerate(vals, 1):
            c = ws_inc.cell(i, j, v)
            c.font = Font(name="Arial", size=8)
            c.border = BDR
            c.fill = PatternFill("solid", fgColor=grav_fill.get(inc.gravidade.value, CINZA))
    
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(settings.EXPORT_DIR, f"BASE_CONSOLIDADA_{ts}.xlsx")
    wb.save(path)
    return path

def gerar_relatorio_pdf(db: Session, competencia_id: int, empresa_nome: str, comp_desc: str) -> str:
    os.makedirs(settings.EXPORT_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(settings.EXPORT_DIR, f"RELATORIO_GERENCIAL_{ts}.pdf")
    
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18*mm, rightMargin=18*mm,
                             topMargin=20*mm, bottomMargin=16*mm)
    AZUL_C = colors.HexColor("#1E3A5F")
    story = [
        Paragraph(f"Relatório Gerencial — {comp_desc}", ParagraphStyle("T", fontSize=18, textColor=AZUL_C, fontName="Helvetica-Bold")),
        Spacer(1, 4*mm),
        Paragraph(f"Empresa: {empresa_nome} · Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
                  ParagraphStyle("S", fontSize=10, textColor=colors.HexColor("#6B7280"))),
        HRFlowable(width="100%", thickness=2, color=colors.HexColor("#2563EB"), spaceAfter=6),
        Spacer(1, 4*mm),
    ]
    
    incs = db.query(Inconsistencia).filter(Inconsistencia.competencia_id == competencia_id).all()
    if incs:
        story.append(Paragraph("Inconsistências", ParagraphStyle("SE", fontSize=12, textColor=AZUL_C, fontName="Helvetica-Bold", spaceBefore=10)))
        data = [[Paragraph(h, ParagraphStyle("H", fontSize=8, textColor=colors.white, fontName="Helvetica-Bold"))
                 for h in ["Módulo","Descrição","Gravidade","Status"]]]
        for inc in incs:
            data.append([inc.modulo, inc.descricao[:60], inc.gravidade.value, inc.status.value])
        t = Table(data, colWidths=[28*mm, 100*mm, 22*mm, 24*mm])
        t.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0), AZUL_C),
            ("FONTSIZE",(0,0),(-1,-1),8),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white, colors.HexColor("#F9FAFB")]),
            ("GRID",(0,0),(-1,-1),0.3, colors.HexColor("#E5E7EB")),
            ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
            ("LEFTPADDING",(0,0),(-1,-1),5),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ]))
        story.append(t)
    
    doc.build(story)
    return path
